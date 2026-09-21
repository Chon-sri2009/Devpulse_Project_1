from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(r"D:\VS_Project\MiniProject_Everything_1")
OUT = ROOT / "deliverables"
SHOT = ROOT / ".artifact-build" / "screenshots"
OUT.mkdir(parents=True, exist_ok=True)

NAVY = "111C34"
BLUE = "3976ED"
GREEN = "17965A"
AMBER = "D4861A"
INK = "172033"
MUTED = "5F6F8C"
PALE = "F3F6FB"
LINE = "D8E1EE"
WHITE = "FFFFFF"
SOFT_BLUE = "EAF1FF"
SOFT_GREEN = "E8F6EE"


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_inches):
    width = Inches(width_inches)
    cell.width = width
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LINE, size=6):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:color"), color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("DevPulse Version 1   |   ")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, value, end])


def set_run_font(run, name="Arial", size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def configure_document(doc, title, subject):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.3)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, color in (("Title", 34, NAVY), ("Heading 1", 22, NAVY), ("Heading 2", 15, BLUE), ("Heading 3", 11, GREEN)):
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    if "Code Block" not in styles:
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = styles["Code Block"]
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    code.font.size = Pt(8.5)
    code.font.color.rgb = RGBColor.from_string(INK)
    code.paragraph_format.left_indent = Inches(0.16)
    code.paragraph_format.right_indent = Inches(0.16)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(6)

    props = doc.core_properties
    props.title = title
    props.subject = subject
    props.author = "DevPulse project team"
    props.keywords = "DevPulse, .NET 9, Blazor Server, diagnostics, deployment"
    props.comments = "Version 1 public project documentation"

    footer = section.footer.paragraphs[0]
    add_page_number(footer)


def add_label(doc, text, color=BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text.upper())
    set_run_font(r, size=9, bold=True, color=color)
    r.font.all_caps = True


def add_cover(doc, title, subtitle, audience, image_name=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("DEVPULSE")
    set_run_font(r, size=13, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(title)
    set_run_font(r, size=34, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run(subtitle)
    set_run_font(r, size=17, bold=False, color=MUTED)

    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Inches(1.55)
    table.columns[1].width = Inches(5.15)
    labels = [("Release", "Version 1"), ("Platform", ".NET 9 Blazor Server"), ("Audience", audience), ("Status", "Presentation and deployment package")]
    first = table.rows[0]
    for idx, (left, right) in enumerate(labels):
        row = first if idx == 0 else table.add_row()
        row.cells[0].text = left
        row.cells[1].text = right
    for row in table.rows:
        set_cell_width(row.cells[0], 1.55)
        set_cell_width(row.cells[1], 5.15)
        set_cell_shading(row.cells[0], NAVY)
        set_cell_shading(row.cells[1], PALE)
        for j, cell in enumerate(row.cells):
            set_cell_margins(cell, 90, 120, 90, 120)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in cell.paragraphs[0].runs:
                set_run_font(run, size=9.5, bold=(j == 0), color=WHITE if j == 0 else INK)
    set_table_borders(table)

    if image_name and (SHOT / image_name).exists():
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        inline = run.add_picture(str(SHOT / image_name), width=Inches(6.65))
        inline._inline.docPr.set("descr", "DevPulse Version 1 workspace overview")
        inline._inline.docPr.set("title", "DevPulse workspace overview")
        p.paragraph_format.space_after = Pt(2)
        cap = doc.add_paragraph("DevPulse Version 1 workspace")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(0)
        for run in cap.runs:
            set_run_font(run, size=8, color=MUTED)
    else:
        doc.add_paragraph()
        p = doc.add_paragraph("Observe. Test. Inspect. Plan. Focus.")
        p.paragraph_format.space_before = Pt(42)
        for run in p.runs:
            set_run_font(run, size=20, bold=True, color=BLUE)


def page(doc, title, label=None, intro=None):
    doc.add_page_break()
    if label:
        add_label(doc, label)
    doc.add_heading(title, level=1)
    if intro:
        p = doc.add_paragraph(intro)
        p.paragraph_format.space_after = Pt(9)
        for run in p.runs:
            set_run_font(run, size=11, color=MUTED)


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.left_indent = Inches(0.24 + level * 0.2)
        p.paragraph_format.first_line_indent = Inches(-0.12)
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_numbered(doc, items):
    for number, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.42)
        p.paragraph_format.first_line_indent = Inches(-0.3)
        p.paragraph_format.space_after = Pt(4)
        n = p.add_run(f"{number}.  ")
        set_run_font(n, size=10, bold=True, color=BLUE)
        p.add_run(item)


def add_code(doc, text):
    p = doc.add_paragraph(style="Code Block")
    p.paragraph_format.keep_together = True
    p.add_run(text)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "EEF2F7")
    p_pr.append(shd)


def add_note(doc, heading, text, kind="info"):
    colors = {"info": (SOFT_BLUE, BLUE), "success": (SOFT_GREEN, GREEN), "warning": ("FFF5E6", AMBER)}
    fill, accent = colors[kind]
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Inches(6.7)
    cell = table.cell(0, 0)
    set_cell_width(cell, 6.7)
    set_cell_shading(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(heading + "\n")
    set_run_font(r, size=10, bold=True, color=accent)
    r = p.add_run(text)
    set_run_font(r, size=9.5, color=INK)
    set_cell_margins(cell, 100, 140, 100, 140)
    set_table_borders(table, accent, 8)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    if widths:
        for i, width in enumerate(widths):
            table.columns[i].width = Inches(width)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, value in enumerate(headers):
        hdr.cells[i].text = value
        if widths:
            set_cell_width(hdr.cells[i], widths[i])
        set_cell_shading(hdr.cells[i], NAVY)
        set_cell_margins(hdr.cells[i])
        for run in hdr.cells[i].paragraphs[0].runs:
            set_run_font(run, size=8.5, bold=True, color=WHITE)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        for i, value in enumerate(values):
            row.cells[i].text = str(value)
            if widths:
                set_cell_width(row.cells[i], widths[i])
            set_cell_shading(row.cells[i], PALE if row_index % 2 == 0 else WHITE)
            set_cell_margins(row.cells[i])
            row.cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            for run in row.cells[i].paragraphs[0].runs:
                set_run_font(run, size=8.3, color=INK)
    set_table_borders(table)
    return table


def add_image(doc, name, caption, width=6.65):
    path = SHOT / name
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    inline = run.add_picture(str(path), width=Inches(width))
    inline._inline.docPr.set("descr", caption)
    inline._inline.docPr.set("title", caption)
    p.paragraph_format.space_after = Pt(2)
    c = doc.add_paragraph(caption)
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(6)
    for run in c.runs:
        set_run_font(run, size=8, color=MUTED)


def add_credits_page(doc):
    doc.add_page_break()
    add_label(doc, "Credits")
    doc.add_heading("Project creators", level=1)
    p = doc.add_paragraph("Add the names of the DevPulse student team and project advisor before sharing or presenting this document.")
    p.paragraph_format.space_after = Pt(14)
    for run in p.runs:
        set_run_font(run, size=11, color=MUTED)
    add_table(doc, ["Role", "Name"], [
        ["Student 1", "________________________________________"],
        ["Student 2", "________________________________________"],
        ["Student 3", "________________________________________"],
        ["Advisor", "________________________________________"],
    ], [1.8, 4.9])
    doc.add_heading("Where to edit the names", level=2)
    doc.add_paragraph("Open page 2 in Microsoft Word. Click the blank line in the Name column, delete the underscores, and type the correct name. Repeat this for all three students and the advisor.")


def build_setup():
    doc = Document()
    configure_document(doc, "DevPulse Version 1 Setup and Deployment Guide", "Installation, configuration, security, and deployment instructions")
    add_cover(doc, "Setup and Deployment Guide", "Install, configure, validate, and publish DevPulse Version 1 safely.", "Developers, operators, reviewers, and project maintainers")
    add_credits_page(doc)

    page(doc, "Quick start", "Start here", "A clean local launch needs only the .NET 9 SDK. Spotify, browser audits, alerts, and private-network checks are optional integrations.")
    doc.add_heading("Five-minute local launch", level=2)
    add_numbered(doc, [
        "Open a terminal in the repository root.",
        "Restore and build the application.",
        "Run the HTTP launch profile.",
        "Open http://127.0.0.1:5214 in a browser.",
        "Confirm that /healthz returns a healthy result.",
    ])
    add_code(doc, "dotnet restore\ndotnet build\ndotnet run --launch-profile http")
    add_note(doc, "What works without credentials", "Overview, network calculators, local inspectors, Focus Lounge wellness tools, and the configured diagnostics workspace can run without Spotify credentials. Production diagnostics are disabled unless the owner explicitly enables them.", "success")
    doc.add_heading("Recommended first demonstration", level=2)
    add_bullets(doc, ["Open Overview and explain that metrics describe the DevPulse host.", "Calculate 192.168.10.25/24 in Network Calculator.", "Run a Website Audit against a site you own.", "Finish in Focus Lounge and optionally connect Spotify."])

    page(doc, "Requirements and repository", "Prepare", "Install the components that match the capabilities you plan to demonstrate.")
    add_table(doc, ["Component", "Required", "Purpose"], [
        [".NET 9 SDK", "Yes", "Build and run the Blazor Server application"],
        ["Visual Studio 2022 17.12+", "Optional", "IDE support for .NET 9"],
        ["Git", "Recommended", "Source control and deployment integration"],
        ["Node.js plus Chromium", "Optional locally", "Browser audit and Lighthouse-style checks"],
        ["Docker", "Optional", "Container validation and hosted deployment"],
        ["Spotify Premium", "Optional", "Spotify Connect controls and browser playback"],
    ], [2.0, 1.0, 3.7])
    doc.add_heading("Important project files", level=2)
    add_table(doc, ["File or folder", "Purpose"], [
        ["MiniProject_Everything_1.csproj", ".NET 9 target and package references"],
        ["Program.cs", "Services, authentication, routes, security middleware"],
        ["appsettings.json", "Safe defaults and non-secret configuration"],
        ["Components/", "Razor pages, layout, and interface components"],
        ["Service/", "Diagnostics, Spotify, networking, audit, and storage logic"],
        ["scripts/", "Browser audit and verification automation"],
        ["Dockerfile", "Production image on port 10000"],
    ], [2.3, 4.4])
    add_note(doc, "Do not change the target framework", "DevPulse Version 1 targets net9.0. Install .NET 9 and keep the project on .NET 9 unless a planned migration is tested separately.", "warning")

    page(doc, "Configuration and secret handling", "Configure", "Keep operational settings readable, but keep every password, token, client secret, and private endpoint outside Git.")
    doc.add_heading("Configuration precedence", level=2)
    add_numbered(doc, ["appsettings.json provides safe committed defaults.", "appsettings.Development.json can provide non-secret development overrides.", ".NET User Secrets provide local credentials.", "Environment variables or a mounted secret file provide hosted credentials."])
    doc.add_heading("Local secrets", level=2)
    add_code(doc, "dotnet user-secrets set \"Spotify:ClientId\" \"YOUR_CLIENT_ID\"\ndotnet user-secrets set \"Spotify:ClientSecret\" \"YOUR_CLIENT_SECRET\"\ndotnet user-secrets set \"Admin:Password\" \"A-unique-password-of-at-least-12-characters\"")
    doc.add_heading("Hosted environment variables", level=2)
    add_code(doc, "Spotify__ClientId=YOUR_CLIENT_ID\nSpotify__ClientSecret=YOUR_CLIENT_SECRET\nAdmin__Password=YOUR_LONG_UNIQUE_PASSWORD\nDiagnostics__Enabled=false")
    add_note(doc, "Leak prevention", "Never place real credentials in appsettings.json, source code, screenshots, test files, Git history, or presentation material. Rotate a credential immediately if it has ever been committed.", "warning")
    doc.add_heading("Persistent cryptographic data", level=2)
    add_bullets(doc, ["Mount a persistent data directory at /app/data in production.", "Keep data-protection keys stable across container replacement.", "Spotify token sessions are encrypted on the server; the browser receives only an opaque session identifier."])

    page(doc, "Spotify integration", "Integrate", "Spotify is optional. The rest of DevPulse remains usable when the integration is not configured.")
    doc.add_heading("Spotify developer dashboard", level=2)
    add_numbered(doc, [
        "Create or open your application in the Spotify Developer Dashboard.",
        "Add your account and any testers allowed by the application mode.",
        "Register the exact redirect URI. Spotify requires an exact character-for-character match.",
        "Store the client ID and client secret using User Secrets or hosted secret variables.",
        "Restart DevPulse, open Focus Lounge, and connect Spotify.",
    ])
    add_table(doc, ["Environment", "Redirect URI"], [
        ["Local HTTP profile", "http://127.0.0.1:5214/signin-spotify"],
        ["Render production", "https://devpulse-project-sng7.onrender.com/signin-spotify"],
        ["Custom domain", "https://YOUR-DOMAIN/signin-spotify"],
    ], [2.0, 4.7])
    add_note(doc, "Redirect mismatch", "Update both the Spotify dashboard and the deployed application URL. Do not add or remove a trailing slash, change localhost to 127.0.0.1, or switch HTTP and HTTPS unless the registered value changes too.", "warning")
    doc.add_heading("Playback behavior", level=2)
    add_bullets(doc, ["Spotify Connect follows the active phone, desktop, speaker, or browser device.", "Selecting Continue on this device transfers playback while preserving play or pause state.", "Browser audio requires Premium and a user gesture to enable browser playback.", "Progress reconciles with Spotify about every four seconds because Spotify provides no playback webhook."])

    page(doc, "Diagnostics, allowlists, and private targets", "Secure", "Network tools run from the DevPulse host, not from the visitor's browser. Owner-controlled allowlists define what the host may contact.")
    doc.add_heading("Enable diagnostics intentionally", level=2)
    add_code(doc, "Diagnostics__Enabled=true\nDiagnostics__FolderRoot=D:\\approved\\logs")
    add_table(doc, ["Setting", "Example", "Effect"], [
        ["Operations__ApprovedHosts__0", "127.0.0.1", "Allows ping, DNS, TLS, and TCP checks for one host"],
        ["Operations__ApprovedPorts__0", "443", "Adds an allowed TCP port"],
        ["Operations__ApprovedUrls__0", "https://api.github.com", "Adds a URL suggestion or exact trusted exception"],
        ["Operations__ApprovedLogFiles__0", "/var/log/app.log", "Allows the log tailer to read one file"],
        ["Operations__Databases__0__Kind", "TCP or Redis", "Selects TCP accept or unauthenticated Redis PING"],
    ], [2.25, 1.8, 2.65])
    add_note(doc, "Private addresses", "Local, private, link-local, and reserved addresses are blocked for arbitrary input. Add an exact owner-approved host when you intentionally test your own LAN or local service. A Render container still cannot reach a private computer behind your router.", "warning")
    doc.add_heading("Bounded execution", level=2)
    add_bullets(doc, ["Load tests are capped at 25 requests and five concurrent workers by default.", "Redirects are disabled for diagnostic HTTP clients.", "Response previews, request bodies, timeouts, scan ports, and crawl resources have explicit limits."])

    page(doc, "Website, API, DNS, and alert setup", "Extend", "These modules are useful for demonstrating DevPulse against systems you own without changing source code.")
    doc.add_heading("Website Audit", level=2)
    add_code(doc, "WebsiteAudit__DefaultUrl=https://example.com/\nWebsiteAudit__MonitoredUrls__0=https://example.com/\nWebsiteAudit__BrowserEnabled=true")
    add_bullets(doc, ["Use a public site you own or are authorized to test.", "The server performs crawl, header, JSON, accessibility, visual, and optional browser checks.", "Scheduled uptime checks run at the configured interval and create incidents when state changes."])
    doc.add_heading("API Collection Runner", level=2)
    add_bullets(doc, ["GET and HEAD may target validated public endpoints.", "POST, PUT, PATCH, and DELETE require the exact endpoint in Operations:ApprovedUrls.", "Authorization values and custom headers remain temporary and are not written to disk."])
    doc.add_heading("DNS and Email Inspector", level=2)
    add_bullets(doc, ["Inspect A, AAAA, MX, SPF, DMARC, DKIM, and CAA records for public domains.", "Use selectors from your mail provider when testing DKIM.", "Results are diagnostic evidence, not a guarantee of message delivery."])
    doc.add_heading("Alerts", level=2)
    add_code(doc, "Alerts__MemoryMb=1024\nAlerts__DiskUsedPercent=90\nAlerts__RequestLatencyMs=2000\nAlerts__Smtp__Password=STORE_AS_A_SECRET")

    page(doc, "Local validation", "Verify", "Run the same checks before presenting or deploying. A successful page load alone is not enough release evidence.")
    doc.add_heading("Build and format", level=2)
    add_code(doc, "dotnet restore\ndotnet build -c Release\ndotnet format --verify-no-changes")
    doc.add_heading("Regression and smoke checks", level=2)
    add_code(doc, "dotnet run --project tests/DevPulse.Tests/DevPulse.Tests.csproj\npowershell -ExecutionPolicy Bypass -File scripts/http-smoke.ps1")
    add_note(doc, "Version 1 verification baseline", "The final Version 1 validation recorded 79 of 79 automated checks passing, a Release build with zero warnings and zero errors, a successful formatting check, and a successful published HTTP smoke suite.", "success")
    doc.add_heading("Manual presentation checks", level=2)
    add_bullets(doc, ["Open every navigation item at desktop width.", "Run one approved network operation and one website audit.", "Upload only non-sensitive sample files to the inspectors.", "Connect Spotify, browse created and saved collections, and confirm the active device behavior.", "Confirm administrator login and logout without enabling process termination."])

    page(doc, "Docker and Render deployment", "Deploy", "The repository Dockerfile publishes a .NET 9 application, installs the browser-audit runtime, and runs as the non-root app user.")
    doc.add_heading("Container contract", level=2)
    add_table(doc, ["Item", "Value"], [
        ["Internal HTTP port", "10000"],
        ["Health endpoint", "/healthz"],
        ["Persistent data path", "/app/data"],
        ["Runtime user", "app (non-root)"],
        ["Browser runtime", "Node.js 24 plus Chromium"],
    ], [2.4, 4.3])
    doc.add_heading("Render setup", level=2)
    add_numbered(doc, [
        "Create a Web Service from the repository and choose Docker as the runtime.",
        "Use port 10000 and configure the health check path as /healthz.",
        "Add Spotify and administrator credentials as secret environment variables.",
        "Mount persistent storage at /app/data if Spotify sessions and telemetry must survive replacement.",
        "Register the final HTTPS Spotify callback and configure trusted proxy addresses or networks.",
        "Deploy one instance unless shared session and telemetry storage are introduced.",
    ])
    add_note(doc, "Reverse proxy safety", "Trust only the proxy ranges supplied by the hosting platform. Do not configure forwarded headers to trust every network, because OAuth depends on the external HTTPS scheme being trustworthy.", "warning")

    page(doc, "Production security checklist", "Protect", "DevPulse contains diagnostics that can reveal host and network information. A public demo needs stricter access controls than a local classroom demonstration.")
    add_table(doc, ["Control", "Required action"], [
        ["Diagnostic exposure", "Keep disabled or place behind authenticated access"],
        ["Administrator access", "Use a unique password of at least 12 characters"],
        ["Process termination", "Keep disabled unless explicitly reviewed"],
        ["Private target checks", "Allowlist only systems you own or are permitted to test"],
        ["Secrets", "Use environment variables or a mounted secret file"],
        ["Data volume", "Persist and protect encryption keys and application state"],
        ["Public reputation", "Resolve Search Console and Safe Browsing warnings before promotion"],
        ["API screenshots", "Never capture tokens, keys, or authorization headers"],
    ], [2.2, 4.5])
    add_note(doc, "Current public launch checkpoint", "Before presenting the public Render URL broadly, complete the Google Search Console or Safe Browsing review and place diagnostic routes behind authenticated access or a private gateway.", "warning")
    doc.add_heading("Safe default", level=2)
    add_code(doc, "Diagnostics__Enabled=false\nAdmin__AllowProcessTermination=false")

    page(doc, "Troubleshooting", "Support", "Start with the symptom, confirm which machine is running the check, and then review the relevant configuration boundary.")
    add_table(doc, ["Symptom", "Likely cause", "Resolution"], [
        ["redirect_uri mismatch", "Registered callback differs from the request", "Make the Spotify dashboard URI exactly match scheme, host, port, path, and slash"],
        ["Diagnostics are disabled", "Production default is off", "Enable Diagnostics only for an intentionally protected deployment"],
        ["Storage shows Render disks", "Checks run on the server host", "Use a local deployment or future remote agent for workstation storage"],
        ["Private IP is blocked", "SSRF protections reject private input", "Add an exact approved host and run DevPulse where that network is reachable"],
        ["No Spotify sound", "Playback remains on another device or browser audio is not activated", "Select an active device or explicitly enable browser playback with Premium"],
        ["Created playlist is missing", "OAuth scopes are stale or the item is not owned/collaborative", "Disconnect, reconnect, and verify playlist ownership and scopes"],
        ["Chrome warns about the site", "Safe Browsing reputation or diagnostic content", "Stop broad sharing, review Search Console, secure routes, and request review"],
        ["Browser audit unavailable", "Node or Chromium path is missing", "Install the optional runtime or use the Docker image"],
    ], [2.05, 2.15, 2.5])

    page(doc, "Release and handoff checklist", "Finish", "Use this page before a presentation, deployment, or handoff to another maintainer.")
    add_bullets(doc, [
        "The repository builds with the .NET 9 SDK.",
        "No real credential is present in committed files or screenshots.",
        "The public hostname and Spotify callback match exactly.",
        "The /healthz endpoint is healthy after deployment.",
        "Persistent storage is mounted when encrypted sessions or telemetry must survive restart.",
        "Diagnostics are disabled or protected by authenticated access.",
        "Only authorized hosts, ports, URLs, log files, and database endpoints are allowlisted.",
        "The Release build, format check, regression executable, and HTTP smoke suite pass.",
        "A manual walkthrough has been completed on the device used for the presentation.",
        "Google Safe Browsing or Search Console status is clear before public promotion.",
    ])
    doc.add_heading("Useful endpoints", level=2)
    add_table(doc, ["Path", "Purpose"], [["/", "Overview"], ["/healthz", "Platform health check"], ["/focus-lounge", "Focus tools and Spotify"], ["/website-audit", "Website assurance"], ["/api-runner", "Bounded API collections"], ["/admin", "Restricted administration"]], [2.0, 4.7])
    add_note(doc, "Version 1 is ready to present", "The application and release package are technically validated. The remaining public-launch work is operational: protect diagnostics and complete the browser reputation review.", "success")

    path = OUT / "DevPulse_V1_Setup_and_Deployment_Guide.docx"
    doc.save(path)
    return path


def build_description():
    doc = Document()
    configure_document(doc, "DevPulse Version 1 Project Description", "Product overview, architecture, capabilities, safeguards, evidence, and roadmap")
    add_cover(doc, "Project Description", "A unified workspace for IT diagnostics, website assurance, network planning, API inspection, and developer focus.", "Technical and nontechnical presentation audiences", "overview.png")
    add_credits_page(doc)

    page(doc, "Executive summary", "Overview", "DevPulse Version 1 is a .NET 9 Blazor Server application that brings common IT and developer checks into one controlled web workspace.")
    doc.add_heading("The problem", level=2)
    doc.add_paragraph("Developers and small operations teams often move between operating-system utilities, network calculators, command-line probes, API clients, website-audit tools, credential-sensitive dashboards, and music applications. That switching makes demonstrations harder and evidence more fragmented.")
    doc.add_heading("The response", level=2)
    doc.add_paragraph("DevPulse combines those workflows without pretending to be an enterprise monitoring replacement. It emphasizes bounded execution, clear ownership, visible safety limits, locally understandable results, and a user interface that can be presented to both technical and nontechnical audiences.")
    add_note(doc, "Version 1 value proposition", "One web workspace can observe its host, test approved services, inspect local data, plan networks, audit owned websites, exercise APIs, and support focused work.", "success")
    doc.add_heading("Release outcome", level=2)
    add_bullets(doc, ["A coherent navigation model across diagnostics, assurance, planning, inspection, and focus tools.", "A server-side safety model for outbound requests and administrative actions.", "A Docker deployment contract for Render or another container host.", "Automated regression, build, formatting, and published HTTP smoke evidence."])

    page(doc, "Users and core workflows", "Audience", "DevPulse is designed for learning, demonstration, small-team operations, and personal development environments.")
    add_table(doc, ["User", "Need", "DevPulse workflow"], [
        ["Developer", "Check an API or website while coding", "API Runner, Website Audit, telemetry, Focus Lounge"],
        ["IT student", "Learn subnets and diagnostic behavior", "Network Calculator, ping, DNS, TLS, TCP scans"],
        ["Small-team operator", "See host and service condition", "Overview, Operations, telemetry, alerts"],
        ["Presenter or reviewer", "Show evidence and system boundaries", "Dashboard, audit results, release status, administration export"],
        ["Music-enabled user", "Stay focused without losing the active Spotify device", "Focus Lounge, Spotify Connect, optional browser playback"],
    ], [1.45, 2.2, 3.2])
    doc.add_heading("Typical end-to-end scenario", level=2)
    add_numbered(doc, ["Confirm application and host status on Overview.", "Calculate the target subnet without manual binary arithmetic.", "Run approved connectivity checks from the DevPulse host.", "Audit an owned website and review browser, header, and accessibility evidence.", "Exercise a public API collection with bounded assertions.", "Use Focus Lounge while Spotify continues on the selected phone or computer."])

    page(doc, "Version 1 capability map", "Product", "The product is organized around five verbs that are easy to explain during a presentation.")
    add_table(doc, ["Pillar", "Modules", "Representative outputs"], [
        ["Observe", "Overview, Telemetry, Administration", "Host metrics, storage, processes, request traces, incidents"],
        ["Test", "Operations, Website Audit, API Runner, DNS & Email", "Latency, status, TLS, crawl results, assertions, DNS policy"],
        ["Inspect", "JSON, JWT, SHA-256, encrypted files", "Trees, decoded claims, hashes, reversible encrypted packages"],
        ["Plan", "Network Calculator", "IPv4/IPv6 ranges, VLSM, routes, bandwidth, MTU and MSS"],
        ["Focus", "Focus Lounge and Spotify", "Timers, breathing, ambience, notes, library and playback controls"],
    ], [1.05, 2.35, 3.45])
    doc.add_heading("Notable feature depth", level=2)
    add_bullets(doc, ["Website auditing includes crawl limits, headers, JSON checks, optional browser evidence, accessibility, screenshots, visual baselines, and scheduled uptime.", "Spotify includes current playback, active-device following, search, Liked Songs, albums, playlists, queue, shuffle, repeat, transfer, and opt-in browser audio.", "Networking includes both active checks and offline calculations, with explicit restrictions for private and reserved targets."])

    page(doc, "Architecture and execution model", "Design", "Blazor Server keeps application logic and credentials on the server while the browser receives an interactive UI over a live circuit.")
    add_table(doc, ["Layer", "Responsibility"], [
        ["Browser", "Renders Razor components, captures user actions, maintains local focus notes, and hosts optional Spotify browser audio"],
        ["Blazor application", "Coordinates pages, authentication, validation, diagnostics, and near-real-time UI updates"],
        ["Application services", "Perform system, storage, network, website, DNS, API, Spotify, telemetry, and administration work"],
        ["Local data directory", "Stores encrypted Spotify sessions, data-protection keys, telemetry history, and selected baselines"],
        ["Approved external services", "Public websites, APIs, DNS, Spotify, alert destinations, and optional OTLP collector"],
    ], [1.65, 5.15])
    add_note(doc, "Critical interpretation", "Every host diagnostic describes the machine running DevPulse. A Render deployment reports its container storage, processes, network path, and reachability. It cannot inspect a visitor's computer without a future authenticated remote agent.", "warning")
    doc.add_heading("Deployment flow", level=2)
    add_code(doc, "Local source -> .NET 9 publish -> Docker image -> HTTPS hosting service\n                                      -> persistent /app/data")
    doc.add_heading("Technology stack", level=2)
    add_bullets(doc, ["ASP.NET Core and Blazor Server on net9.0.", "OpenTelemetry instrumentation with optional OTLP export.", "Node.js and Chromium for optional browser audits.", "Docker multi-stage image running as a non-root user."])

    page(doc, "Safety and trust model", "Security", "Diagnostics are useful only when the application prevents them from becoming an unrestricted proxy, scanner, credential sink, or destructive control surface.")
    add_table(doc, ["Boundary", "Version 1 control"], [
        ["Outbound targets", "Public-target validation, pinned addresses, exact owner allowlists, redirect blocking"],
        ["Request volume", "Timeout, response-size, request-count, concurrency, port, and crawl caps"],
        ["Credentials", "User Secrets or environment variables; masked transient API authorization values"],
        ["Spotify sessions", "Encrypted server-side tokens, opaque browser cookie, seven-day maximum session"],
        ["Administration", "Separate administrator cookie, antiforgery protection, rate limiting, audit history"],
        ["Process termination", "Disabled by default and requires explicit owner enablement"],
        ["Production routes", "Diagnostic surfaces can return 404 when Diagnostics is disabled"],
    ], [2.0, 4.8])
    add_note(doc, "Deployment recommendation", "Public deployments should put diagnostic pages behind authenticated access or a private gateway. Spotify authentication is separate and does not grant administrator rights.", "warning")
    doc.add_heading("Data minimization", level=2)
    add_bullets(doc, ["Uploaded inspector files are processed in memory rather than retained as user content.", "API collections exist only in the active Blazor circuit.", "Query strings and authorization data are excluded or redacted from telemetry and exports."])

    page(doc, "Website and API assurance", "Evidence", "DevPulse turns a public site or API that the user owns into a repeatable demonstration target.")
    add_image(doc, "website-audit.png", "Website Audit combines bounded server checks with optional browser evidence.", 6.25)
    doc.add_heading("Website Audit outputs", level=2)
    add_bullets(doc, ["HTTP status, redirects, timing, headers, metadata, resources, and crawl findings.", "Optional browser run with screenshot, accessibility evidence, performance signals, and visual comparison.", "Scheduled availability checks that create telemetry incidents when state changes."])
    doc.add_heading("API and DNS outputs", level=2)
    add_bullets(doc, ["Up to ten temporary API requests with methods, headers, bodies, and assertions.", "Public DNS inspection for address, mail, SPF, DMARC, DKIM, and CAA records.", "Strict write-method and private-address boundaries to preserve owner intent."])

    page(doc, "Developer wellness and Spotify", "Experience", "Focus Lounge keeps restorative tools and music control beside technical work without forcing playback into the browser.")
    add_image(doc, "focus-lounge.png", "Focus Lounge combines a timer, breathing, generated ambience, reflection, and Spotify.", 6.25)
    add_table(doc, ["Experience", "Behavior"], [
        ["Focus timer", "Alternates focused work and break periods while the tab remains open"],
        ["Box breathing", "Guides four equal phases and can be stopped at any time"],
        ["Generated ambience", "Produces rain, ocean, and brown-noise sound without an external playlist"],
        ["Private reflection", "Keeps the note in browser-local storage"],
        ["Spotify Connect", "Follows the active desktop, phone, speaker, or browser device"],
        ["Browser playback", "Requires Premium and explicit activation by the user"],
    ], [2.0, 4.8])

    page(doc, "Quality and release evidence", "Verification", "Version 1 was validated at service, integration, formatting, publishing, and user-interface levels.")
    add_table(doc, ["Check", "Result", "Coverage"], [
        ["Automated regression executable", "79 / 79 passed", "Spotify parsing and sessions, cryptography, safety rules, audit behavior, telemetry, APIs, DNS"],
        ["Release build", "0 warnings, 0 errors", "Complete net9.0 solution build"],
        ["Formatting", "Verified", "Repository source formatting"],
        ["Published HTTP smoke suite", "Passed", "Routes, static assets, diagnostics gate, OAuth/PKCE, scopes, token endpoint, proxy handling, logout"],
        ["Visual review", "Passed", "Desktop navigation and presentation-critical pages"],
    ], [2.0, 1.35, 3.45])
    add_note(doc, "Release statement", "The software package is technically ready for a controlled presentation and deployment. Public promotion still depends on securing diagnostic routes and resolving the current Google Safe Browsing or Search Console warning.", "success")
    doc.add_heading("Demonstration evidence", level=2)
    add_bullets(doc, ["Overview provides immediate system context.", "Network Calculator produces deterministic examples without contacting a target.", "Website Audit demonstrates permission-aware external assurance.", "Focus Lounge provides a clear, nontechnical closing experience."])

    page(doc, "Known limitations and release notes", "Honesty", "These boundaries should be stated during a presentation so that expectations match what Version 1 actually does.")
    add_table(doc, ["Limitation", "Practical meaning"], [
        ["Server-side visibility", "Hosted metrics and storage belong to the Render container, not the visitor"],
        ["Private network reachability", "A hosted container cannot normally reach a home or campus LAN"],
        ["Spotify near-real-time polling", "External device changes reconcile periodically because there is no playback webhook"],
        ["Spotify account constraints", "Premium is required for playback controls and browser audio; development apps may restrict users"],
        ["Database health depth", "Version 1 performs TCP acceptance or unauthenticated Redis PING, not SQL query analysis"],
        ["Single-instance storage", "Local encrypted sessions and journals are not shared across replicas"],
        ["Public reputation review", "The current public hostname must clear browser reputation review before broad sharing"],
    ], [2.2, 4.6])
    add_note(doc, "Why these limits are useful", "They keep Version 1 understandable and safer while establishing clear work for a future remote agent, shared storage, database monitoring, and role-based access.", "info")

    page(doc, "Presentation route and Version 2", "Next steps", "A short live path shows the product clearly without depending on every external integration.")
    doc.add_heading("Suggested live demonstration", level=2)
    add_numbered(doc, [
        "Overview: explain host identity and the difference between local and hosted metrics.",
        "Network Calculator: calculate 192.168.10.25/24 and point out the usable range.",
        "Website Audit: run or show cached evidence for a website you own.",
        "API Runner: perform one public GET request and explain write-method allowlists.",
        "Focus Lounge: start the timer, demonstrate breathing, and show Spotify device continuity.",
    ])
    doc.add_heading("Version 2 candidates", level=2)
    add_table(doc, ["Candidate", "Value"], [
        ["SQL Server health and performance monitor", "Adds query latency, connection, wait, storage, and backup visibility"],
        ["IT asset inventory and IP address management", "Tracks devices, owners, addresses, subnets, and lifecycle"],
        ["Authenticated remote monitoring agent", "Collects workstation or server metrics beyond the hosted container"],
        ["Incident response and backup verification", "Adds runbooks, evidence, recovery checks, and notification workflows"],
        ["Role-based access control", "Separates viewers, operators, auditors, and administrators"],
    ], [2.7, 4.1])
    add_note(doc, "Version 1 conclusion", "DevPulse Version 1 establishes the product language, navigation, safety model, deployment contract, and quality baseline needed for these additions.", "success")

    path = OUT / "DevPulse_V1_Project_Description.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    print(build_setup())
    print(build_description())
